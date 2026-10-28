"""Per-kind parsers (13.3). Each returns a list of Parsed documents for one file.

A Parsed document carries its fields, its chunks (text + provenance), and for
transcripts its timed segments; song_bank.json also yields song rows. Nothing
here guesses: an unknown transcript shape is reported, not interpreted.
"""
from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field

CHUNK_CHARS = 1500
WINDOW_SEC = 60.0
VIDEO_ID_RE = re.compile(r"^([A-Za-z0-9_-]{11})(?=$|[_.\- ])")
ISO_DATE_RE = re.compile(r"(20\d\d)[-_](\d\d)[-_](\d\d)")
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
MON_DAY_RE = re.compile(r"(?<![a-z])(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*[_-]?(\d{1,2})(?!\d)", re.I)
BRACKET_RE = re.compile(r"^\[(\d{1,2}):(\d{2})(?::(\d{2}))?\]\s*(.*)$")
VTT_TIME_RE = re.compile(r"^(?:(\d{1,2}):)?(\d{2}):(\d{2})[.,](\d{1,3})\s+-->\s+")
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic", ".ico"}
VIDEO_EXT = {".mp4", ".mov", ".m4v", ".mkv", ".webm", ".avi"}
AUDIO_EXT = {".wav", ".mp3", ".m4a", ".aac", ".flac", ".ogg"}
CODE_EXT = {".py", ".bat", ".sh", ".ps1", ".cmd", ".js", ".mjs", ".ts", ".css", ".html", ".htm", ".jsx", ".tsx"}
TEXT_EXT = {".md", ".txt", ".log", ".yaml", ".yml", ".toml", ".ini", ".cfg"}
PRIVATE_PARTS = ("emails/", "memory/", "cloud_workspace/mem/")


@dataclass
class Parsed:
    kind: str
    title: str = ""
    video_id: str | None = None
    date: str | None = None
    private: bool = False
    history: bool = False
    extra: dict = field(default_factory=dict)
    chunks: list = field(default_factory=list)      # dicts: text, start_line, end_line, start_sec, end_sec, link
    segments: list = field(default_factory=list)    # dicts: start, dur, text
    songs: list = field(default_factory=list)


class Unknown(Exception):
    """A shape the parser will not guess at. The message goes to build.log."""


# ---------- small helpers -------------------------------------------------------
def looks_like_video_id(tok: str) -> bool:
    if len(tok) != 11 or not re.fullmatch(r"[A-Za-z0-9_-]{11}", tok):
        return False
    if any(c.isdigit() for c in tok) or "-" in tok or "_" in tok:
        return True
    return sum(c.isupper() for c in tok) >= 2 and any(c.islower() for c in tok)


def video_id_from_name(path: str) -> str | None:
    base = path.rsplit("/", 1)[-1]
    m = VIDEO_ID_RE.match(base)
    return m.group(1) if m and looks_like_video_id(m.group(1)) else None


def date_from(path: str, *values) -> str | None:
    for v in list(values) + [path]:
        if not v:
            continue
        m = ISO_DATE_RE.search(str(v))
        if m:
            return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    return None


def date_hint(path: str) -> str | None:
    """Month and day without a year ('oct18'), kept as a hint, never turned into a date."""
    m = MON_DAY_RE.search(path.rsplit("/", 1)[-1].lower())
    if m and 1 <= int(m.group(2)) <= 31:
        return f"{m.group(1).lower()}{int(m.group(2))}"
    return None


def is_private(path: str) -> bool:
    p = "/" + path.lower()
    return any("/" + part in p for part in PRIVATE_PARTS)


def hms(t) -> float | None:
    """'1:02:03', '62:03', '0:00', 12.3 -> seconds."""
    if t is None:
        return None
    if isinstance(t, (int, float)):
        return float(t)
    parts = str(t).strip().split(":")
    try:
        nums = [float(x) for x in parts]
    except ValueError:
        return None
    s = 0.0
    for n in nums:
        s = s * 60 + n
    return s


def yt_link(video_id: str | None, sec: float | None) -> str | None:
    if not video_id or sec is None:
        return None
    return f"https://youtu.be/{video_id}?t={int(sec)}"


def words_of_path(path: str) -> str:
    s = re.sub(r"([a-z])([A-Z])", r"\1 \2", path)
    return " ".join(w for w in re.split(r"[/_\-.\s]+", s) if w)


def chunk_lines(lines: list[str], title: str = "", heading_aware: bool = False) -> list[dict]:
    """~1,500-character chunks, by heading then paragraph, each with its line range (1-based)."""
    out, buf, start = [], [], None

    def flush(end):
        nonlocal buf, start
        text = "\n".join(buf).strip()
        if text:
            out.append({"text": text, "start_line": start, "end_line": end})
        buf, start = [], None

    for i, ln in enumerate(lines, 1):
        is_head = heading_aware and ln.startswith("#")
        if (is_head and buf) or (buf and sum(len(x) + 1 for x in buf) + len(ln) > CHUNK_CHARS and not ln.strip()):
            flush(i - 1)
        if buf and sum(len(x) + 1 for x in buf) > CHUNK_CHARS * 1.6:
            flush(i - 1)
        if start is None:
            start = i
        buf.append(ln)
    flush(len(lines))
    return out


def window(segments: list[dict], video_id: str | None) -> list[dict]:
    """Timed segments -> ~60-second chunks with a youtu.be timestamp link."""
    out, buf, w0 = [], [], None
    for s in segments:
        if w0 is None:
            w0 = s["start"]
        if s["start"] - w0 >= WINDOW_SEC and buf:
            out.append(_win(buf, video_id))
            buf, w0 = [], s["start"]
        buf.append(s)
    if buf:
        out.append(_win(buf, video_id))
    return out


def _win(buf, vid):
    start = buf[0]["start"]
    end = max(x["start"] + (x.get("dur") or 0) for x in buf)
    return {"text": " ".join(x["text"].strip() for x in buf if x["text"]), "start_sec": start, "end_sec": end, "link": yt_link(vid, start)}


def fill_durations(segs: list[dict]) -> list[dict]:
    segs = [s for s in segs if s.get("start") is not None and (s.get("text") or "").strip()]
    segs.sort(key=lambda s: s["start"])
    for i, s in enumerate(segs):
        if s.get("dur") in (None, 0):
            nxt = segs[i + 1]["start"] if i + 1 < len(segs) else None
            s["dur"] = round(nxt - s["start"], 3) if nxt is not None and nxt > s["start"] else 0.0
    return segs


# ---------- transcripts ---------------------------------------------------------
def transcript_json(path: str, data: bytes) -> list[Parsed]:
    d = json.loads(data.decode("utf-8-sig"))
    segs, shape = None, None
    if isinstance(d, dict) and isinstance(d.get("segments"), list):
        shape = "segments"
        segs = [{"start": hms(s.get("start")), "dur": s.get("dur"), "text": s.get("text") or ""} for s in d["segments"] if isinstance(s, dict)]
    elif isinstance(d, dict) and isinstance(d.get("cues"), list) and d["cues"] and isinstance(d["cues"][0], dict) and "t" in d["cues"][0]:
        shape = "cues t/s"
        segs = [{"start": hms(c.get("t")), "dur": None, "text": c.get("s") or ""} for c in d["cues"]]
    elif isinstance(d, dict) and isinstance(d.get("cues"), list) and d["cues"] and isinstance(d["cues"][0], dict) and "start" in d["cues"][0]:
        shape = "cues start/text"
        segs = [{"start": hms(c.get("start")), "dur": c.get("dur"), "text": c.get("text") or ""} for c in d["cues"]]
    elif isinstance(d, list) and d and isinstance(d[0], dict) and "start" in d[0] and "text" in d[0]:
        shape = "list start/text"
        segs = [{"start": hms(c.get("start")), "dur": c.get("dur"), "text": c.get("text") or ""} for c in d]
    if segs is None:
        keys = sorted(d.keys()) if isinstance(d, dict) else [type(d).__name__]
        raise Unknown(f"unknown transcript shape (keys: {', '.join(map(str, keys))[:200]})")
    vid = (d.get("video_id") if isinstance(d, dict) else None) or video_id_from_name(path)
    segs = fill_durations(segs)
    p = Parsed(kind="transcript", title=(d.get("title") if isinstance(d, dict) else "") or f"Transcript {vid or path}",
               video_id=vid, date=date_from(path, d.get("published") if isinstance(d, dict) else None,
                                            d.get("service_date") if isinstance(d, dict) else None),
               private=is_private(path),
               extra={"source": (d.get("source") if isinstance(d, dict) else None) or path.rsplit("/", 1)[-1], "shape": shape,
                      "date_hint": date_hint(path)})
    p.segments = segs
    p.chunks = window(segs, vid)
    return [p]


def transcript_vtt(path: str, data: bytes) -> list[Parsed]:
    text = data.decode("utf-8-sig", "replace")
    segs, cur = [], None
    for ln in text.splitlines():
        m = VTT_TIME_RE.match(ln.strip())
        if m:
            h, mi, s, ms = m.groups()
            cur = {"start": int(h or 0) * 3600 + int(mi) * 60 + int(s) + int(ms.ljust(3, "0")) / 1000, "dur": None, "text": ""}
            segs.append(cur)
        elif cur is not None and ln.strip() and not ln.strip().isdigit():
            cur["text"] = (cur["text"] + " " + re.sub(r"<[^>]+>", "", ln.strip())).strip()
        elif not ln.strip():
            cur = None
    vid = video_id_from_name(path)
    segs = fill_durations(segs)
    p = Parsed(kind="transcript", title=f"Transcript {vid or path.rsplit('/', 1)[-1]}", video_id=vid, date=date_from(path),
               private=is_private(path), extra={"source": path.rsplit("/", 1)[-1], "shape": "vtt", "date_hint": date_hint(path)})
    p.segments, p.chunks = segs, window(segs, vid)
    return [p]


def transcript_text(path: str, data: bytes) -> list[Parsed]:
    """*_para.txt and work/chap/*_transcript.txt: lines starting [h:mm:ss], or plain paragraphs."""
    lines = data.decode("utf-8-sig", "replace").splitlines()
    segs, cur = [], None
    for ln in lines:
        m = BRACKET_RE.match(ln.strip())
        if m:
            a, b, c, rest = m.groups()
            sec = int(a) * 3600 + int(b) * 60 + int(c) if c is not None else int(a) * 60 + int(b)
            cur = {"start": float(sec), "dur": None, "text": rest}
            segs.append(cur)
        elif cur is not None and ln.strip():
            cur["text"] += " " + ln.strip()
    vid = video_id_from_name(path)
    p = Parsed(kind="transcript", title=f"Transcript {vid or path.rsplit('/', 1)[-1]}", video_id=vid, date=date_from(path),
               private=is_private(path), extra={"source": path.rsplit("/", 1)[-1], "shape": "bracketed" if segs else "plain",
                                                "date_hint": date_hint(path)})
    if segs:
        p.segments = fill_durations(segs)
        p.chunks = window(p.segments, vid)
    else:
        p.chunks = chunk_lines(lines)
    return [p]


def chapters(path: str, data: bytes) -> list[Parsed]:
    d = json.loads(data.decode("utf-8-sig"))
    items = d.get("chapters") if isinstance(d, dict) else d
    if not isinstance(items, list):
        raise Unknown("chapters file without a chapters list")
    vid = (d.get("video_id") if isinstance(d, dict) else None) or video_id_from_name(path)
    p = Parsed(kind="chapters", title=f"Chapters {vid or path}", video_id=vid, date=date_from(path), private=is_private(path))
    for c in items:
        if not isinstance(c, dict):
            continue
        t = hms(c.get("t") or c.get("time") or c.get("start"))
        label = c.get("label") or c.get("title") or ""
        p.chunks.append({"text": f"{c.get('t') or ''} {label}".strip(), "start_sec": t, "end_sec": t, "link": yt_link(vid, t)})
    return [p]


# ---------- songs, setlists, speakers ----------------------------------------------
def norm_title(t: str) -> str:
    t = (t or "").lower().replace("’", "'").replace("‘", "'").replace("“", '"').replace("”", '"')
    t = re.sub(r"[^\w\s']", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def song_bank(path: str, data: bytes) -> list[Parsed]:
    d = json.loads(data.decode("utf-8-sig"))
    rows = d.items() if isinstance(d, dict) else [(x.get("title"), x) for x in d if isinstance(x, dict)]
    out = []
    for title, v in rows:
        if not isinstance(v, dict) or str(title).startswith("_"):
            continue
        ccli = str(v["ccli"]).strip() if v.get("ccli") not in (None, "") else None
        p = Parsed(kind="song", title=str(title))
        fields = {k: v.get(k) for k in ("ccli", "writers", "copyright", "admin", "credit_line", "original", "verified", "note", "performed_by")}
        p.extra = {k: x for k, x in fields.items() if x not in (None, "")}
        p.chunks = [{"text": "\n".join(f"{k}: {x}" for k, x in [("title", title)] + list(p.extra.items()))}]
        p.songs = [{"title": str(title), "title_norm": norm_title(str(title)), "ccli": ccli, "writers": v.get("writers"),
                    "copyright": v.get("copyright"), "admin": v.get("admin"), "credit_line": v.get("credit_line"),
                    "original": 1 if v.get("original") else 0, "verified": v.get("verified"), "note": v.get("note")}]
        out.append(p)
    return out


def setlist(path: str, data: bytes) -> list[Parsed]:
    d = json.loads(data.decode("utf-8-sig"))
    if not isinstance(d, dict) or not isinstance(d.get("songs"), list):
        raise Unknown("setlist without a songs list")
    date = date_from(path, d.get("date_ct"), d.get("date"))
    lines = []
    for i, s in enumerate(d["songs"], 1):
        if isinstance(s, dict):
            bits = [s.get("title") or "?"]
            if s.get("ccli"):
                bits.append(f"CCLI {s['ccli']}")
            if s.get("writer") or s.get("writers"):
                bits.append(str(s.get("writer") or s.get("writers")))
            lines.append(f"{i}. " + " · ".join(bits))
        else:
            lines.append(f"{i}. {s}")
    p = Parsed(kind="setlist", title=f"Setlist {date or ''} {d.get('service') or ''}".strip(), date=date,
               extra={"service": d.get("service"), "source": d.get("source")})
    p.chunks = [{"text": f"{d.get('service') or ''} {date or ''}\n" + "\n".join(lines)}]
    return [p]


def speakers(path: str, data: bytes) -> list[Parsed]:
    d = json.loads(data.decode("utf-8-sig"))
    out = []
    for group, items in (d.items() if isinstance(d, dict) else []):
        if group.startswith("_") or not isinstance(items, list):
            continue
        for it in items:
            if not isinstance(it, dict):
                continue
            name = it.get("name") or it.get("heard_as") or "?"
            p = Parsed(kind="speaker", title=str(name), extra={"group": group, **{k: v for k, v in it.items() if k != "name"}})
            p.chunks = [{"text": f"{group}: " + "; ".join(f"{k}: {v}" for k, v in it.items())}]
            out.append(p)
    if isinstance(d, dict) and d.get("_comment"):
        p = Parsed(kind="speaker", title="speakers.json: the rules", extra={"group": "_comment"})
        p.chunks = [{"text": "\n".join(d["_comment"]) if isinstance(d["_comment"], list) else str(d["_comment"])}]
        out.append(p)
    return out


# ---------- descriptions and history --------------------------------------------------
def _video_records(d) -> list[tuple[str, dict]]:
    """(video_id, record) pairs from the shapes the toolkit writes: a dict keyed by video id,
    a list of objects with video_id/id, YouTube API objects (id + snippet)."""
    out = []
    if isinstance(d, dict):
        for k, v in d.items():
            if looks_like_video_id(str(k)) and isinstance(v, (dict, list)):
                out.append((k, v if isinstance(v, dict) else {"entries": v}))
            elif isinstance(v, dict) and looks_like_video_id(str(v.get("video_id") or v.get("id") or "")):
                out.append((v.get("video_id") or v.get("id"), v))
        if not out:
            for key in ("videos", "items", "entries"):
                if isinstance(d.get(key), list):
                    return _video_records(d[key])
    elif isinstance(d, list):
        for v in d:
            if isinstance(v, dict):
                vid = v.get("video_id") or v.get("videoId") or v.get("id")
                if isinstance(vid, dict):
                    vid = vid.get("videoId")
                if vid and looks_like_video_id(str(vid)):
                    out.append((vid, v))
    return out


def _record_text(v: dict) -> tuple[str, str, list]:
    sn = v.get("snippet") if isinstance(v.get("snippet"), dict) else {}
    title = v.get("title") or sn.get("title") or ""
    desc = v.get("description") or sn.get("description") or v.get("preview") or ""
    tags = v.get("tags") or sn.get("tags") or []
    if not desc and isinstance(v.get("entries"), list):
        desc = "\n".join(json.dumps(e, ensure_ascii=False)[:600] for e in v["entries"][:10])
    return str(title), str(desc), tags if isinstance(tags, list) else []


def description_json(path: str, data: bytes, history: bool = False) -> list[Parsed]:
    d = json.loads(data.decode("utf-8-sig"))
    recs = _video_records(d)
    if not recs:
        raise Unknown("no video records")
    out = []
    for vid, v in recs:
        title, desc, tags = _record_text(v)
        sn = v.get("snippet") if isinstance(v.get("snippet"), dict) else {}
        p = Parsed(kind="history" if history else "description", title=title or vid, video_id=vid, history=history,
                   date=date_from(path, v.get("published"), sn.get("publishedAt"), v.get("date")), private=is_private(path),
                   extra={"tags": tags[:40], "source": path.rsplit("/", 1)[-1]} if tags else {"source": path.rsplit("/", 1)[-1]})
        body = "\n".join(x for x in [title, desc, ("tags: " + ", ".join(map(str, tags))) if tags else ""] if x)
        p.chunks = [{"text": body}]
        out.append(p)
    return out


def description_text(path: str, data: bytes) -> list[Parsed]:
    lines = data.decode("utf-8-sig", "replace").splitlines()
    vid = video_id_from_name(path)
    title = next((ln.strip() for ln in lines if ln.strip()), path.rsplit("/", 1)[-1])
    p = Parsed(kind="description", title=title[:200], video_id=vid, date=date_from(path), private=is_private(path),
               extra={"source": path.rsplit("/", 1)[-1], "date_hint": date_hint(path)})
    p.chunks = chunk_lines(lines)
    return [p]


# ---------- notes, tables, code, generic data ------------------------------------------
def note(path: str, data: bytes) -> list[Parsed]:
    lines = data.decode("utf-8-sig", "replace").splitlines()
    md = path.lower().endswith(".md")
    title = next((ln.lstrip("# ").strip() for ln in lines if ln.strip()), path.rsplit("/", 1)[-1]) if md else path.rsplit("/", 1)[-1]
    p = Parsed(kind="note", title=title[:200], video_id=video_id_from_name(path), date=date_from(path), private=is_private(path),
               extra={"date_hint": date_hint(path)} if date_hint(path) else {})
    p.chunks = chunk_lines(lines, heading_aware=md)
    return [p]


def data_json(path: str, data: bytes) -> list[Parsed]:
    """JSON that is not one of the known kinds: its strings, flattened, as a note-like 'data' document."""
    d = json.loads(data.decode("utf-8-sig"))
    lines = []

    def walk(x, key=""):
        if isinstance(x, dict):
            for k, v in x.items():
                walk(v, f"{key}.{k}" if key else str(k))
        elif isinstance(x, list):
            for i, v in enumerate(x[:5000]):
                walk(v, key)
        elif isinstance(x, str) and x.strip():
            lines.append(f"{key}: {x}" if key else x)
        elif isinstance(x, (int, float)) and key:
            lines.append(f"{key}: {x}")

    walk(d)
    p = Parsed(kind="data", title=path.rsplit("/", 1)[-1], video_id=video_id_from_name(path), date=date_from(path),
               private=is_private(path))
    p.chunks = [{"text": c["text"]} for c in chunk_lines(lines)]
    return [p]


def table_csv(path: str, data: bytes) -> list[Parsed]:
    text = data.decode("utf-8-sig", "replace")
    rows = list(csv.reader(io.StringIO(text)))
    p = Parsed(kind="table", title=path.rsplit("/", 1)[-1], date=date_from(path), private=is_private(path))
    if rows:
        head = rows[0]
        for i, r in enumerate(rows[1:], 2):
            cells = "; ".join(f"{h}: {v}" for h, v in zip(head, r) if v.strip())
            if cells:
                p.chunks.append({"text": cells, "start_line": i, "end_line": i})
    return [p]


def table_xlsx(path: str, data: bytes) -> list[Parsed]:
    try:
        import openpyxl
    except ImportError as e:
        raise Unknown("xlsx skipped: openpyxl missing") from e
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    p = Parsed(kind="table", title=path.rsplit("/", 1)[-1], date=date_from(path), private=is_private(path), extra={"sheets": wb.sheetnames})
    for ws in wb.worksheets:
        head = None
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = ["" if v is None else str(v) for v in row]
            if head is None:
                if any(vals):
                    head = vals
                continue
            cells = "; ".join(f"{h or f'col{j}'}: {v}" for j, (h, v) in enumerate(zip(head, vals), 1) if v.strip())
            if cells:
                p.chunks.append({"text": f"[{ws.title}] {cells}", "start_line": i, "end_line": i})
    return [p]


def code(path: str, data: bytes) -> list[Parsed]:
    lines = data.decode("utf-8", "replace").splitlines()
    p = Parsed(kind="code", title=path.rsplit("/", 1)[-1])
    p.chunks = chunk_lines(lines)
    return [p]


# ---------- media -------------------------------------------------------------------
def pptx_text(data: bytes) -> list[str]:
    out = []
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)), key=lambda n: int(re.findall(r"\d+", n)[-1]))
        for n in slides:
            xml = z.read(n).decode("utf-8", "replace")
            out.append(" ".join(re.findall(r"<a:t>([^<]*)</a:t>", xml)))
    return out


def media(path: str, size: int, sha256: str, meta: dict, text_pages: list[str] | None = None) -> list[Parsed]:
    base = path.rsplit("/", 1)[-1]
    p = Parsed(kind="media", title=base, video_id=video_id_from_name(path), date=date_from(path), private=is_private(path),
               extra={k: v for k, v in meta.items() if v not in (None, "")})
    desc = [words_of_path(path), base]
    for k in ("width", "height", "duration", "codec", "format", "note"):
        if meta.get(k):
            desc.append(f"{k} {meta[k]}")
    p.chunks = [{"text": " ".join(desc)}]
    for i, t in enumerate(text_pages or [], 1):
        if t.strip():
            p.chunks.append({"text": t, "start_line": i, "end_line": i})
    return [p]
